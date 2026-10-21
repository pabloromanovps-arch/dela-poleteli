#!/usr/bin/env python3
"""Презентация из deck.yaml в стиле бренда → редактируемый PPTX (+ PDF и превью).

  python tools/deck/build_deck.py examples/deck/concept_deck.yaml --brand studio --out out/deck.pptx --pdf --preview

Типы слайдов: cover, section, statement, text, image, image_grid, compare, palette,
typography, quote, timeline, end. Пути к картинкам — относительно deck.yaml;
отсутствующая картинка → плашка-заглушка с подписью (для черновиков).
"""
import argparse
import glob
import os
import shutil
import subprocess
import sys
import tempfile

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

TOOLS = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, TOOLS)
import brand as B  # noqa: E402

SW, SH = Emu(12192000), Emu(6858000)  # 16:9, 13.333×7.5 in
M = int(SW * 0.06)                     # поля 6%
COL = (SW - 2 * M) / 12                # 12 колонок


class Deck:
    def __init__(self, brand, base_dir):
        self.b = brand
        self.base = base_dir
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SW, SH
        self.blank = self.prs.slide_layouts[6]
        self.head = brand["fonts"]["heading"]["family"]
        self.body = brand["fonts"]["body"]["family"]
        self.n = 0

    def c(self, role):
        if role == "muted":
            return RGBColor(*self.muted())
        return RGBColor(*B.rgb(B.color(self.b, role)))

    @staticmethod
    def _lum(rgb):
        f = lambda v: (v / 255 / 12.92) if v / 255 <= 0.03928 else ((v / 255 + 0.055) / 1.055) ** 2.4
        r, g, b = (f(v) for v in rgb)
        return 0.2126 * r + 0.7152 * g + 0.0722 * b

    def _contrast(self, a, b):
        la, lb = sorted((self._lum(a), self._lum(b)), reverse=True)
        return (la + 0.05) / (lb + 0.05)

    def muted(self):
        """Приглушённый цвет служебного текста: смесь text и bg, контраст к фону ≥ 4,5:1.
        Роль secondary в брендах часто задана как «декор, текст запрещён» — для текста она не годится."""
        if not hasattr(self, "_muted"):
            fg, bg = B.rgb(B.color(self.b, "text")), B.rgb(B.color(self.b, "bg"))
            t = 0.55
            while True:
                mix = tuple(round(bg[i] + (fg[i] - bg[i]) * t) for i in range(3))
                if self._contrast(mix, bg) >= 4.5 or t >= 1:
                    break
                t += 0.05
            self._muted = mix
        return self._muted

    # ---------- примитивы ----------

    def new(self, bg="bg"):
        s = self.prs.slides.add_slide(self.blank)
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = self.c(bg)
        self.n += 1
        return s

    def text(self, s, x, y, w, h, txt, size, role="text", font=None, bold=False, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, spacing=None):
        if role == "secondary":  # текст — только читаемым приглушённым цветом
            role = "muted"
        box = s.shapes.add_textbox(int(x), int(y), int(w), int(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(str(txt).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            if spacing:
                p.line_spacing = spacing
            r = p.add_run()
            r.text = line
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.name = font or self.body
            r.font.color.rgb = self.c(role)
        return box

    def rect(self, s, x, y, w, h, role=None, rgb=None):
        shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(x), int(y), int(w), int(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb or self.c(role)
        shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def picture(self, s, path, x, y, w, h, label="", fit="auto"):
        """fit: cover — заполнить рамку с кадрированием; contain — вписать целиком;
        auto — contain, если пропорции кадра и рамки различаются более чем в 1,6 раза (широкие панорамы)."""
        full = os.path.join(self.base, path) if path else ""
        if path and os.path.exists(full):
            from PIL import Image
            iw, ih = Image.open(full).size
            box_r, img_r = w / h, iw / ih
            if fit == "auto":
                fit = "contain" if max(img_r / box_r, box_r / img_r) > 1.6 else "cover"
            if fit == "contain":
                k = min(w / iw, h / ih)
                pw, ph = iw * k, ih * k
                self.contained = True
                return s.shapes.add_picture(full, int(x + (w - pw) / 2), int(y + (h - ph) / 2), int(pw), int(ph))
            self.contained = False
            pic = s.shapes.add_picture(full, int(x), int(y), int(w), int(h))
            if img_r > box_r:
                crop = (1 - box_r / img_r) / 2
                pic.crop_left = pic.crop_right = crop
            else:
                crop = (1 - img_r / box_r) / 2
                pic.crop_top = pic.crop_bottom = crop
            return pic
        self.rect(s, x, y, w, h, rgb=RGBColor(0x26, 0x26, 0x2B))
        self.text(s, x, y, w, h, f"[ {label or path or 'изображение'} ]", 14, role="secondary",
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

    def footer(self, s, sd):
        if sd.get("no_footer"):
            return
        self.text(s, M, SH - M * 0.75, COL * 8, M * 0.4, self.footer_text, 10, role="secondary")
        self.text(s, SW - M - COL * 2, SH - M * 0.75, COL * 2, M * 0.4, f"{self.n:02d}", 10,
                  role="secondary", align=PP_ALIGN.RIGHT)

    def kicker(self, s, txt):
        if txt:
            self.text(s, M, M, COL * 8, Pt(20), txt.upper(), 12, role="accent", bold=True)

    # ---------- типы слайдов ----------

    def cover(self, sd):
        s = self.new()
        if sd.get("image"):
            self.picture(s, sd["image"], 0, 0, SW, SH)
            self.rect(s, 0, SH * 0.55, SW, SH * 0.45, role="bg").fill.transparency = 0.2
        self.kicker(s, sd.get("kicker"))
        self.text(s, M, SH * 0.58, COL * 10, SH * 0.25, sd["title"], 60, font=self.head, bold=True,
                  anchor=MSO_ANCHOR.BOTTOM, spacing=0.9)
        if sd.get("subtitle"):
            self.text(s, M, SH * 0.85, COL * 10, SH * 0.06, sd["subtitle"], 18, role="secondary")

    def section(self, sd):
        s = self.new()
        # номер и заголовок — по общей нижней линии
        self.text(s, M, SH * 0.3, COL * 2, SH * 0.25, sd.get("number", ""), 60, role="accent",
                  font=self.head, bold=True, anchor=MSO_ANCHOR.BOTTOM)
        self.text(s, M + COL * 2, SH * 0.3, COL * 10, SH * 0.25, sd["title"], 44, font=self.head, bold=True,
                  anchor=MSO_ANCHOR.BOTTOM)
        self.footer(s, sd)

    def statement(self, sd):
        s = self.new()
        self.kicker(s, sd.get("kicker"))
        self.text(s, M, SH * 0.2, COL * 11, SH * 0.6, sd["title"], 54, font=self.head, bold=True,
                  anchor=MSO_ANCHOR.MIDDLE, spacing=0.95)
        if sd.get("note"):
            self.text(s, M, SH * 0.78, COL * 7, SH * 0.1, sd["note"], 16, role="secondary")
        self.footer(s, sd)

    def text_slide(self, sd):
        s = self.new()
        self.kicker(s, sd.get("kicker"))
        self.text(s, M, M + Pt(36), COL * 5, SH * 0.5, sd["title"], 32, font=self.head, bold=True, spacing=0.95)
        body = sd.get("body", "")
        if isinstance(body, list):
            body = "\n".join(f"—  {b}" for b in body)
        self.text(s, M + COL * 6, M + Pt(40), COL * 6, SH * 0.7, body, 18, spacing=1.2)
        self.footer(s, sd)

    def image(self, sd):
        s = self.new()
        if sd.get("full", True):
            self.contained = False
            self.picture(s, sd.get("image"), 0, 0, SW, SH, sd.get("caption"), sd.get("fit", "auto"))
            if sd.get("caption"):
                if self.contained:  # панорама вписана целиком — подпись под ней, на фоне слайда
                    self.kicker(s, sd.get("kicker"))
                    self.text(s, M, SH - M - Pt(40), COL * 11, Pt(40), sd["caption"], 14)
                else:
                    self.rect(s, M - Pt(8), SH - M - Pt(30), COL * 6, Pt(30), role="bg").fill.transparency = 0.15
                    self.text(s, M, SH - M - Pt(26), COL * 6, Pt(24), sd["caption"], 14)
        else:
            self.kicker(s, sd.get("kicker"))
            self.picture(s, sd.get("image"), M, M + Pt(36), SW - 2 * M, SH - 2 * M - Pt(70), sd.get("caption"))
            self.text(s, M, SH - M - Pt(20), COL * 10, Pt(20), sd.get("caption", ""), 14, role="secondary")
            self.footer(s, {"no_footer": True})

    def image_grid(self, sd):
        s = self.new()
        self.kicker(s, sd.get("kicker"))
        imgs = sd["images"]
        cols = sd.get("cols", 3 if len(imgs) > 4 else 2)
        rows = (len(imgs) + cols - 1) // cols
        top = M + Pt(36)
        gap = Pt(10)
        cap_h = Pt(22)
        cw = (SW - 2 * M - gap * (cols - 1)) / cols
        ch = (SH - top - M - Pt(20) - (gap + cap_h) * rows) / rows
        for i, it in enumerate(imgs):
            it = it if isinstance(it, dict) else {"image": it}
            x = M + (i % cols) * (cw + gap)
            y = top + (i // cols) * (ch + gap + cap_h)
            self.picture(s, it.get("image"), x, y, cw, ch, os.path.basename(it.get("image") or ""))
            if it.get("caption"):
                self.text(s, x, y + ch + Pt(4), cw, cap_h, it["caption"], 11, role="secondary")
        self.footer(s, sd)

    def compare(self, sd):
        s = self.new()
        self.kicker(s, sd.get("kicker"))
        self.text(s, M, M + Pt(30), COL * 12, Pt(50), sd["title"], 32, font=self.head, bold=True)
        opts = sd["options"]
        n = len(opts)
        gap = Pt(16)
        w = (SW - 2 * M - gap * (n - 1)) / n
        top = M + Pt(110)
        for i, o in enumerate(opts):
            x = M + i * (w + gap)
            rec = o.get("recommended")
            self.rect(s, x, top, w, Pt(4), role="accent" if rec else "secondary")
            self.text(s, x, top + Pt(16), w, Pt(40), o["name"] + ("  ★" if rec else ""), 22, font=self.head,
                      bold=True, role="accent" if rec else "text")
            pts = o.get("points", [])
            self.text(s, x, top + Pt(64), w, SH - top - M - Pt(80), "\n".join(f"—  {p}" for p in pts), 15,
                      spacing=1.2)
        self.footer(s, sd)

    def palette(self, sd):
        s = self.new()
        self.kicker(s, sd.get("kicker", "Палитра"))
        self.text(s, M, M + Pt(30), COL * 12, Pt(50), sd.get("title", ""), 32, font=self.head, bold=True)
        cols = sd.get("colors")
        if not cols:  # роли бренда; одинаковые цвета объединяются в одну плашку
            merged = {}
            for k, v in self.b["colors"].items():
                merged.setdefault(v["hex"].upper(), []).append(k)
            share = self.b.get("proportions", {})
            cols = [{"name": " / ".join(n), "hex": h, "share": sum(share.get(x, 0) for x in n) or None}
                    for h, n in merged.items()]
        n = len(cols)
        gap = Pt(10)
        w = (SW - 2 * M - gap * (n - 1)) / n
        top = M + Pt(110)
        h = SH - top - M - Pt(70)
        for i, c in enumerate(cols):
            x = M + i * (w + gap)
            shp = self.rect(s, x, top, w, h, rgb=RGBColor(*B.rgb(c["hex"])))
            shp.line.color.rgb = self.c("secondary")
            shp.line.width = Pt(0.75)
            self.text(s, x, top + h + Pt(10), w, Pt(20), c["name"], 14, bold=True)
            self.text(s, x, top + h + Pt(32), w, Pt(20), c["hex"].upper() + (f"  {c['share']}%" if c.get("share") else ""),
                      12, role="secondary")
        self.footer(s, sd)

    def typography(self, sd):
        s = self.new()
        self.kicker(s, sd.get("kicker", "Типографика"))
        y = M + Pt(40)
        for lvl in sd.get("levels") or [
            {"name": "Заголовок", "size": 54, "font": self.head, "bold": True, "sample": "Дела полетели"},
            {"name": "Подзаголовок", "size": 28, "font": self.head, "bold": True, "sample": "Свет, ритм и масштаб"},
            {"name": "Текст", "size": 18, "font": self.body, "sample": "Съешь же ещё этих мягких французских булок, да выпей чаю."},
        ]:
            self.text(s, M, y, COL * 3, Pt(20), f"{lvl['name']} · {lvl.get('font', self.body)} · {lvl['size']} pt", 11,
                      role="secondary")
            self.text(s, M + COL * 3, y - Pt(4), COL * 9, Pt(lvl["size"] * 1.3), lvl["sample"], lvl["size"],
                      font=lvl.get("font", self.body), bold=lvl.get("bold", False))
            y += Pt(lvl["size"] * 1.3 + 36)
        self.footer(s, sd)

    def quote(self, sd):
        s = self.new()
        self.text(s, M, SH * 0.18, Pt(80), Pt(80), "«", 96, role="accent", font=self.head, bold=True)
        self.text(s, M + COL, SH * 0.25, COL * 10, SH * 0.45, sd["text"], 36, font=self.head, spacing=1.05)
        if sd.get("author"):
            self.text(s, M + COL, SH * 0.75, COL * 8, Pt(24), "— " + sd["author"], 16, role="secondary")
        self.footer(s, sd)

    def timeline(self, sd):
        s = self.new()
        self.kicker(s, sd.get("kicker", "План"))
        self.text(s, M, M + Pt(30), COL * 12, Pt(50), sd.get("title", ""), 32, font=self.head, bold=True)
        steps = sd["steps"]
        n = len(steps)
        w = (SW - 2 * M) / n
        y = SH * 0.5
        self.rect(s, M, y, SW - 2 * M, Pt(2), role="secondary")
        for i, st in enumerate(steps):
            x = M + i * w
            self.rect(s, x, y - Pt(5), Pt(12), Pt(12), role="accent")
            self.text(s, x, y - Pt(50), w - Pt(10), Pt(30), st.get("when", ""), 14, role="accent", bold=True)
            self.text(s, x, y + Pt(22), w - Pt(14), SH * 0.3, st["what"], 15, spacing=1.15)
        self.footer(s, sd)

    def end(self, sd):
        s = self.new()
        self.text(s, M, SH * 0.3, COL * 10, SH * 0.25, sd.get("title", "Следующий шаг"), 54, font=self.head, bold=True)
        if sd.get("body"):
            self.text(s, M, SH * 0.58, COL * 8, SH * 0.2, sd["body"], 20, role="secondary", spacing=1.2)
        if sd.get("contact"):
            self.text(s, M, SH - M - Pt(24), COL * 10, Pt(24), sd["contact"], 14, role="accent")


def to_pdf(pptx_path):
    out_dir = os.path.dirname(os.path.abspath(pptx_path))
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        print("LibreOffice не найден — PDF пропущен (PowerPoint/Keynote: Экспорт в PDF)")
        return None
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", out_dir, pptx_path],
                   check=True, capture_output=True)
    return os.path.splitext(os.path.abspath(pptx_path))[0] + ".pdf"


def preview(pdf_path):
    out = os.path.splitext(pdf_path)[0] + "_preview.png"
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["pdftoppm", "-png", "-r", "40", pdf_path, os.path.join(tmp, "s")], check=True)
        subprocess.run([sys.executable, os.path.join(TOOLS, "contact_sheet.py"), tmp, "--cols", "4",
                        "--thumb", "480", "--out", out], check=True)
    return out


def main():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("deck")
    p.add_argument("--brand", default="studio")
    p.add_argument("--out", required=True)
    p.add_argument("--pdf", action="store_true")
    p.add_argument("--preview", action="store_true", help="контактный лист слайдов (требует --pdf)")
    a = p.parse_args()

    with open(a.deck, encoding="utf-8") as f:
        spec = yaml.safe_load(f)
    d = Deck(B.load(spec.get("brand", a.brand)), os.path.dirname(os.path.abspath(a.deck)))
    d.footer_text = spec.get("footer", spec.get("title", ""))
    kinds = {"cover": d.cover, "section": d.section, "statement": d.statement, "text": d.text_slide,
             "image": d.image, "image_grid": d.image_grid, "compare": d.compare, "palette": d.palette,
             "typography": d.typography, "quote": d.quote, "timeline": d.timeline, "end": d.end}
    for i, sd in enumerate(spec["slides"], 1):
        if sd["type"] not in kinds:
            sys.exit(f"Слайд {i}: неизвестный тип {sd['type']}")
        kinds[sd["type"]](sd)
        if sd.get("notes"):
            d.prs.slides[-1].notes_slide.notes_text_frame.text = sd["notes"]

    os.makedirs(os.path.dirname(os.path.abspath(a.out)), exist_ok=True)
    d.prs.save(a.out)
    print(f"PPTX: {a.out} ({d.n} слайдов)")
    if a.pdf or a.preview:
        pdf = to_pdf(a.out)
        if pdf:
            print(f"PDF: {pdf}")
            if a.preview:
                print(f"Превью: {preview(pdf)}")


if __name__ == "__main__":
    main()
